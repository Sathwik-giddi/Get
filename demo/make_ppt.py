from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = Path.home() / "Desktop" / "HackSprint_Round2_Get.pptx"

SW, SH = Inches(13.333), Inches(7.5)

CREAM = RGBColor(0xF7, 0xF3, 0xEC)
TEAL = RGBColor(0x45, 0xB3, 0xB3)
TEAL_DARK = RGBColor(0x0B, 0x5E, 0x5E)
TEAL_PALE = RGBColor(0xD9, 0xF0, 0xF0)
TEAL_MID = RGBColor(0x7F, 0xD2, 0xD2)
INK = RGBColor(0x23, 0x27, 0x2E)
MUTED = RGBColor(0x5F, 0x5B, 0x52)
HAIRLINE = RGBColor(0xD8, 0xD2, 0xC4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SERIF = "Georgia"
SANS = "Calibri"

LEFT = Inches(0.7)
CONTENT_W = Inches(13.333 - 1.4)
TITLE_Y = Inches(0.35)


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    return prs


def add_slide(prs: Presentation, with_wash: bool = True):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = CREAM
    if with_wash:
        wash = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(8.9), Inches(0), Inches(4.433), SH
        )
        wash.line.fill.background()
        fill = wash.fill
        fill.gradient()
        fill.gradient_angle = 0.0
        stops = fill.gradient_stops
        stops[0].position = 0.0
        stops[0].color.rgb = TEAL_PALE
        stops[1].position = 1.0
        stops[1].color.rgb = TEAL_MID
    return slide


def textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    box.text_frame.vertical_anchor = anchor
    return box


def para(frame, text, size=18, bold=False, color=INK, font=SANS, align=PP_ALIGN.LEFT,
         space_after=Pt(4), space_before=Pt(0)):
    p = frame.add_paragraph() if frame.paragraphs else frame.paragraphs[0]
    if frame.paragraphs and len(frame.paragraphs) > 1 or (frame.paragraphs and frame.paragraphs[0].runs):
        p = frame.add_paragraph()
    p.alignment = align
    p.space_after = space_after
    p.space_before = space_before
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return p


def title(slide, text, sub=None, width=None, size=46):
    box = textbox(slide, LEFT, TITLE_Y, CONTENT_W if width is None else width, Inches(1.4))
    frame = box.text_frame
    frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = TEAL
    run.font.name = SERIF
    if sub:
        para(frame, sub, size=17, color=MUTED, space_before=Pt(2))
    return box


def diagram_box(slide, left, top, width, height, lines, accent=False, center=True,
                fill=WHITE, border=HAIRLINE, title_size=13, body_size=11):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = TEAL_DARK if accent else border
    shape.line.width = Pt(2.25 if accent else 1.25)
    shape.text_frame.word_wrap = True
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    shape.text_frame.margin_left = Inches(0.12)
    shape.text_frame.margin_right = Inches(0.12)
    shape.text_frame.margin_top = Inches(0.06)
    shape.text_frame.margin_bottom = Inches(0.06)
    frame = shape.text_frame
    first = True
    for text, size, bold, color in lines:
        if first:
            p = frame.paragraphs[0]
            first = False
        else:
            p = frame.add_paragraph()
        p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
        p.space_after = Pt(2)
        p.space_before = Pt(0)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(title_size if size == "t" else body_size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = SANS
    return shape


def arrow(slide, x1, y1, x2, y2, color=TEAL_DARK, width=Pt(2)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = width
    return conn


def badge(slide, cx, cy, text, r=Inches(0.17), fill=TEAL_DARK, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - r, cy - r, r * 2, r * 2)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.text_frame.word_wrap = True
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = color
    run.font.name = SANS
    return shape


def bullets(slide, left, top, width, items, size=15, gap=Pt(8)):
    box = textbox(slide, left, top, width, Inches(4))
    frame = box.text_frame
    first = True
    for head, body in items:
        p = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        p.space_after = gap
        p.space_before = Pt(0)
        if head:
            run = p.add_run()
            run.text = head + "  "
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = TEAL_DARK
            run.font.name = SANS
        run = p.add_run()
        run.text = body
        run.font.size = Pt(size)
        run.font.color.rgb = INK
        run.font.name = SANS
    return box


# ---------------------------------------------------------------- slides


def slide_cover(prs):
    slide = add_slide(prs)
    box = textbox(slide, LEFT, Inches(0.55), Inches(7.6), Inches(2.2))
    frame = box.text_frame
    run = frame.paragraphs[0].add_run()
    run.text = "HackSprint"
    run.font.size = Pt(84)
    run.font.bold = True
    run.font.color.rgb = TEAL
    run.font.name = SERIF
    frame.paragraphs[0].alignment = PP_ALIGN.LEFT

    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, LEFT, Inches(2.35), Inches(4.4), Inches(0.62)
    )
    pill.fill.solid()
    pill.fill.fore_color.rgb = TEAL_PALE
    pill.line.fill.background()
    pill.text_frame.word_wrap = True
    pill.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = pill.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "24-Hour Hackathon"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = TEAL_DARK
    run.font.name = SANS

    meta = textbox(slide, LEFT, Inches(3.25), Inches(7.6), Inches(0.9))
    frame = meta.text_frame
    p = frame.paragraphs[0]
    for text, bold in [("17-18 October  |  ", True),
                       ("Manipal Academy of Higher Education (MAHE)", False)]:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(16)
        run.font.bold = bold
        run.font.color.rgb = TEAL_DARK if bold else INK
        run.font.name = SANS

    proj = textbox(slide, LEFT, Inches(4.35), Inches(7.6), Inches(1.3))
    frame = proj.text_frame
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = "Get"
    run.font.size = Pt(54)
    run.font.bold = True
    run.font.color.rgb = INK
    run.font.name = SERIF
    para(frame, "One situation, several sources, one call you can check.",
         size=17, color=MUTED, space_before=Pt(2))

    fields = textbox(slide, LEFT, Inches(5.85), Inches(5.4), Inches(1.1))
    frame = fields.text_frame
    first = True
    for label, value in [("YOUR NAME", "Sathwik Giddi"),
                         ("TEAM NAME", "Zero"),
                         ("TRACK NO", "P42")]:
        p = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        p.space_after = Pt(3)
        run = p.add_run()
        run.text = label + ":   "
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = INK
        run.font.name = SANS
        run = p.add_run()
        run.text = value
        run.font.size = Pt(15)
        run.font.color.rgb = TEAL_DARK
        run.font.name = SANS

    bar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, LEFT, Inches(6.85), Inches(7.6), Inches(0.5)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL_DARK
    bar.line.fill.background()
    bar.text_frame.word_wrap = True
    bar.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = bar.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Team Size 1-4 Members      |      Round 2: Idea Submission"
    run.font.size = Pt(14)
    run.font.color.rgb = WHITE
    run.font.name = SANS


def slide_solution(prs):
    slide = add_slide(prs)
    title(slide, "Solution", "What Get does, in one pass through the product.")
    top = Inches(1.85)
    bw, bh = Inches(3.55), Inches(2.5)
    gap = (CONTENT_W.inches - 3 * bw.inches) / 2
    steps = [
        ("1  FRAGMENTS IN",
         "A dispatch note, a phone photo, a PDF bulletin. Contradictory, partial, real."),
        ("2  ONE SKILL PER SOURCE",
         "The planner assigns analyze_document or analyze_image to each source. Typed JSON out."),
        ("3  ONE CALL YOU CAN CHECK",
         "make_decision reconciles everything: the call, the chain, the evidence, the confidence."),
    ]
    x = LEFT
    for i, (head, body) in enumerate(steps):
        diagram_box(slide, x, top, bw, bh,
                    [(head, "t", True, TEAL_DARK), (body, "b", False, INK)],
                    accent=(i == 2), title_size=15, body_size=13)
        if i < 2:
            badge(slide, x + bw + Inches(gap / 2), top + bh / 2, str(i + 1))
        x = x + bw + Inches(gap)
    bullets(slide, LEFT, Inches(4.75), CONTENT_W, [
        ("Provenance, not prose.",
         "The explanation is the pipeline that ran, not a paragraph written after the answer."),
        ("One decider.",
         "Only make_decision may decide. Every upstream skill extracts and reports."),
        ("Logged by construction.",
         "Every run appends an audit row: JSONL locally, BigQuery when flagged."),
    ], size=15)


def slide_stack(prs):
    slide = add_slide(prs)
    title(slide, "Tech Stack and Architecture", "Everything named here is in the repository.",
          width=Inches(8.0), size=40)
    top = Inches(1.85)
    col = Inches(3.75)
    layers = [
        ("WEB  :3000", "React 19 + Vite\nInputs, call, trace rail, audit table"),
        ("API  :8123", "FastAPI + uvicorn\n/process, /health, /audit, serves the built UI"),
        ("AGENT", "Planner, 3 skills as SKILL.md\nGemini 2.0 Flash, Vertex AI ready"),
    ]
    x = LEFT
    y = top
    for head, body in layers:
        diagram_box(slide, x, y, col, Inches(1.45),
                    [(head, "t", True, TEAL_DARK), (body, "b", False, INK)],
                    title_size=14, body_size=12)
        y = y + Inches(1.45) + Inches(0.28)
    midx = (LEFT + x + col) / 2 if False else None
    cx = LEFT + col / 2
    arrow(slide, cx, top + Inches(1.45), cx, top + Inches(1.73))
    arrow(slide, cx, top + Inches(1.45 + 1.73), cx, top + Inches(1.73 + 1.73))

    rx = LEFT + col + Inches(0.7)
    rw = CONTENT_W - col - Inches(0.7)
    diagram_box(slide, rx, top, rw, Inches(2.2),
                [("CLOUD, BEHIND FLAGS", "t", True, TEAL_DARK),
                 ("Vertex AI via GOOGLE_CLOUD_PROJECT", "b", False, INK),
                 ("Cloud Storage mirror via GCS_BUCKET", "b", False, INK),
                 ("BigQuery audit via GET_BIGQUERY", "b", False, INK)],
                title_size=14, body_size=12)
    diagram_box(slide, rx, top + Inches(2.2) + Inches(0.28), rw, Inches(1.9),
                [("CONTRACT", "t", True, TEAL_DARK),
                 ("Pydantic schemas on every model call", "b", False, INK),
                 ("temperature 0.1, validate then retry", "b", False, INK)],
                title_size=14, body_size=12)
    bullets(slide, rx, Inches(6.35), Inches(8.7) - rx, [
        ("Deploy.", "One Dockerfile: Python deps, npm build, uvicorn serves both. One Cloud Run service."),
    ], size=13)


def slide_dataflow(prs):
    slide = add_slide(prs)
    title(slide, "Data Flow Diagram", "The convoy run, top to bottom of the trace.")
    top = Inches(1.8)
    cols = [
        ("SOURCES", [
            "operator_note\nthe situation",
            "dispatch_log_7kilara.txt\nfield report",
            "bulletin_narmada_rainfall.pdf\nforecast",
            "IMG_4471_bridge.jpg\nphone photo",
        ]),
        ("PLAN", [
            "analyze_document\nbulletin pdf",
            "analyze_document\noperator note",
            "analyze_document\ndispatch txt",
            "analyze_image\nbridge photo",
        ]),
        ("EXECUTE", [
            "DocumentFacts\nfacts, risk flags",
            "DocumentFacts\nintent, window",
            "DocumentFacts\nfield observations",
            "ImageFindings\nseverity: high",
        ]),
        ("DECIDE", [
            "make_decision\nHold Convoy B, reroute via Route 7. Confidence 93%.",
        ]),
        ("AUDIT", [
            "JSONL row now\nBigQuery when flagged",
        ]),
    ]
    n = len(cols)
    cw = CONTENT_W / n
    row_h = Inches(0.92)
    row_gap = Inches(0.14)
    boxes = []
    for ci, (head, items) in enumerate(cols):
        x = LEFT + cw * ci + Inches(0.08)
        w = cw - Inches(0.16)
        hbox = textbox(slide, x, top, w, Inches(0.35))
        p = hbox.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = head
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = WHITE if x > Inches(8.9) else TEAL_DARK
        run.font.name = SANS
        y = top + Inches(0.42)
        col_boxes = []
        for item in items:
            head_line, _, body = item.partition("\n")
            is_decide = ci == 3
            b = diagram_box(slide, x, y, w, row_h if not is_decide else row_h * 2 + row_gap,
                            [(head_line, "t", True, TEAL_DARK if is_decide else INK),
                             (body, "b", False, INK)],
                            accent=is_decide, title_size=12, body_size=10.5)
            col_boxes.append(b)
            y = y + (row_h if not is_decide else row_h * 2 + row_gap) + row_gap
        boxes.append(col_boxes)
    for ci in range(n - 1):
        left_col, right_col = boxes[ci], boxes[ci + 1]
        for li, lbox in enumerate(left_col):
            target = right_col[min(li, len(right_col) - 1)]
            y1 = lbox.top + lbox.height / 2
            y2 = target.top + target.height / 2
            arrow(slide, lbox.left + lbox.width, int(y1),
                  target.left, int(y2), width=Pt(1.5))
    cap = textbox(slide, LEFT, Inches(6.95), Inches(8.0), Inches(0.4))
    p = cap.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = ("4 sources, 5 skill calls, one call. make_decision runs last and always runs.")
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED
    run.font.name = SANS


def slide_screenshot(prs):
    slide = add_slide(prs)
    title(slide, "Screenshot of Your Project", "Live run, convoy scenario. Prototype engine.")
    img_path = ROOT / "demo" / "shots" / "02_convoy.png"
    top = Inches(1.8)
    ih = Inches(4.9)
    iw = ih * (1600 / 1100)
    slide.shapes.add_picture(str(img_path), LEFT, top, width=iw, height=ih)
    rx = LEFT + iw + Inches(0.5)
    rw = CONTENT_W - iw - Inches(0.5)
    bullets(slide, rx, top + Inches(0.1), rw, [
        ("The call.", "One imperative sentence, who acts next, confidence with its reason."),
        ("How it got there.", "Plan and reasoning chain on one numbered rail, sources in the gutter."),
        ("Evidence.", "Attributed per source, marked for, against, or context."),
        ("Prototype output.", "Amber banner states the engine. Gemini key swaps it live."),
    ], size=13)


def slide_others(prs):
    slide = add_slide(prs)
    title(slide, "Others", "Run it, test it, and what comes next.")
    top = Inches(1.8)
    mid = LEFT + CONTENT_W / 2
    bullets(slide, LEFT, top, CONTENT_W / 2 - Inches(0.35), [
        ("Code.", "Repository link goes here. The tree holds agent, api, web, skills, demo, tests."),
        ("Run.", "./run.sh dev starts the API on 8123 and the web UI on 3000."),
        ("Prove.", "34 acceptance checks plus a browser gate that measures contrast, keyboard, viewports."),
    ], size=14)
    bullets(slide, mid + Inches(0.35), top, CONTENT_W / 2 - Inches(0.35), [
        ("Honest today.", "Mock engine until GEMINI_API_KEY is set. The UI says which engine ran."),
        ("Next.", "Live Gemini key, Cloud Run deploy from the Dockerfile, second scenario in the demo."),
        ("Ask.", "Convoy B is waiting on Route 7. Questions welcome."),
    ], size=14)


def main() -> int:
    for name in ["Sathwik Giddi", "Zero", "P42", "Get"]:
        assert name, "cover field missing"
    prs = new_deck()
    slide_cover(prs)
    slide_solution(prs)
    slide_stack(prs)
    slide_dataflow(prs)
    slide_screenshot(prs)
    slide_others(prs)
    for path in [str(ROOT / "demo" / "shots" / "02_convoy.png")]:
        assert Path(path).is_file(), f"missing {path}"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"saved {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())